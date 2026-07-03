"""GA 1그룹 지사 성과지표 군집분석 및 엑셀 보고서 생성 스크립트."""
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
import numpy as np
import pandas as pd
from sklearn.cluster import KMeans
from sklearn.metrics import silhouette_score
from sklearn.preprocessing import StandardScaler

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CSV_PATH = os.path.join(BASE_DIR, "ga_branch_data.csv")
OUT_DIR = os.path.join(BASE_DIR, "output")
XLSX_PATH = os.path.join(OUT_DIR, "GA_1그룹_지사분석_보고서.xlsx")
IMG_DIR = os.path.join(OUT_DIR, "_charts")
os.makedirs(IMG_DIR, exist_ok=True)

KEY_COLS = ["상호명_가상", "사업단_가상", "조직명_가상"]
METRICS = [
    "보장성NCEV_가상",
    "가동5만_가상",
    "가동20만_가상",
    "실수금률_가상",
    "장기위험손해율_가상",
    "장기1차년도손해율_가상",
    "장기2차년도손해율_가상",
    "유지율_4_가상",
    "유지율_7_가상",
    "유지율_13_가상",
    "유지율_25_37_가상",
]

# 한글 폰트 (환경에 나눔고딕 등이 없을 수 있어 가능한 후보 중 존재하는 것 사용)
for cand in ["NanumGothic", "Malgun Gothic", "AppleGothic", "Noto Sans CJK KR"]:
    if any(cand.lower() in f.name.lower() for f in fm.fontManager.ttflist):
        plt.rcParams["font.family"] = cand
        break
plt.rcParams["axes.unicode_minus"] = False

NAVY = "1F3864"
BLUE = "2E75B6"
LIGHT_BLUE = "D9E7F5"
GREY = "7F7F7F"
PALETTE = ["#2E75B6", "#C55A11", "#548235", "#7030A0", "#BF9000"]


def load_group1():
    df = pd.read_csv(CSV_PATH, encoding="cp949")
    g1 = df[df["그룹분류"] == "1그룹"].copy().reset_index(drop=True)
    assert len(g1) == 472, f"1그룹 행 수가 예상과 다릅니다: {len(g1)}"
    assert g1[KEY_COLS].duplicated().sum() == 0, "key 중복 존재"
    assert g1[METRICS].isna().sum().sum() == 0, "결측치 존재"
    return g1


def winsorize(series, lower=0.01, upper=0.99):
    lo, hi = series.quantile(lower), series.quantile(upper)
    return series.clip(lo, hi)


def cluster(g1):
    clipped = g1[METRICS].apply(winsorize)
    X = StandardScaler().fit_transform(clipped.values)

    sil_scores = {}
    for k in range(3, 6):
        km = KMeans(n_clusters=k, random_state=42, n_init=10).fit(X)
        sil_scores[k] = silhouette_score(X, km.labels_)

    best_k = max(sil_scores, key=sil_scores.get)
    km = KMeans(n_clusters=best_k, random_state=42, n_init=10).fit(X)
    g1["군집"] = km.labels_

    sizes = g1["군집"].value_counts().sort_index()
    print("실루엣 스코어:", {k: round(v, 4) for k, v in sil_scores.items()})
    print("선정된 k:", best_k)
    print("군집 크기:", sizes.to_dict())
    assert sizes.min() >= 10, "지나치게 작은(이상치 성) 군집이 존재합니다."
    return g1, best_k, sil_scores


def build_profile(g1):
    overall_mean = g1[METRICS].mean()
    overall_std = g1[METRICS].std()
    cluster_mean = g1.groupby("군집")[METRICS].mean()
    cluster_median = g1.groupby("군집")[METRICS].median()
    zscore = (cluster_mean - overall_mean) / overall_std
    return cluster_mean, cluster_median, zscore


def name_clusters(zscore):
    names = {}
    for cid, row in zscore.iterrows():
        top = row.abs().sort_values(ascending=False).index[:2]
        parts = []
        for m in top:
            direction = "높음" if row[m] > 0 else "낮음"
            label = m.replace("_가상", "").replace("_", " ")
            parts.append(f"{label} {direction}")
        names[cid] = " / ".join(parts)
    return names


def make_charts(g1, cluster_mean, zscore, best_k, sil_scores):
    paths = {}

    # 1) 실루엣 스코어 막대그래프 (k 선정 근거)
    fig, ax = plt.subplots(figsize=(5, 3.2), dpi=150)
    ks = list(sil_scores.keys())
    vals = list(sil_scores.values())
    colors = [PALETTE[0] if k == best_k else "#BFBFBF" for k in ks]
    ax.bar([str(k) for k in ks], vals, color=colors)
    ax.set_title("군집 수(k)별 실루엣 스코어", fontsize=11)
    ax.set_xlabel("k")
    ax.set_ylabel("Silhouette Score")
    fig.tight_layout()
    p = os.path.join(IMG_DIR, "silhouette.png")
    fig.savefig(p)
    plt.close(fig)
    paths["silhouette"] = p

    # 2) 군집별 지사 수
    fig, ax = plt.subplots(figsize=(5, 3.2), dpi=150)
    sizes = g1["군집"].value_counts().sort_index()
    ax.bar([f"군집 {i}" for i in sizes.index], sizes.values, color=PALETTE[: len(sizes)])
    ax.set_title("군집별 지사 수", fontsize=11)
    for i, v in enumerate(sizes.values):
        ax.text(i, v + 3, str(v), ha="center", fontsize=9)
    fig.tight_layout()
    p = os.path.join(IMG_DIR, "cluster_size.png")
    fig.savefig(p)
    plt.close(fig)
    paths["size"] = p

    # 3) z-score 히트맵
    fig, ax = plt.subplots(figsize=(9, 4), dpi=150)
    data = zscore.T.values
    im = ax.imshow(data, cmap="RdBu_r", vmin=-1.5, vmax=1.5, aspect="auto")
    ax.set_xticks(range(zscore.shape[0]))
    ax.set_xticklabels([f"군집 {i}" for i in zscore.index])
    ax.set_yticks(range(len(METRICS)))
    ax.set_yticklabels([m.replace("_가상", "") for m in METRICS], fontsize=8)
    for i in range(data.shape[0]):
        for j in range(data.shape[1]):
            ax.text(j, i, f"{data[i, j]:.2f}", ha="center", va="center", fontsize=7)
    ax.set_title("군집별 지표 z-score (전체 평균 대비)", fontsize=11)
    fig.colorbar(im, ax=ax, shrink=0.8)
    fig.tight_layout()
    p = os.path.join(IMG_DIR, "zscore_heatmap.png")
    fig.savefig(p)
    plt.close(fig)
    paths["heatmap"] = p

    # 4) 군집별 레이더 차트
    labels = [m.replace("_가상", "") for m in METRICS]
    angles = np.linspace(0, 2 * np.pi, len(labels), endpoint=False).tolist()
    angles += angles[:1]
    fig, ax = plt.subplots(figsize=(6.5, 6.5), dpi=150, subplot_kw=dict(polar=True))
    for idx, (cid, row) in enumerate(zscore.iterrows()):
        vals = row.tolist()
        vals += vals[:1]
        ax.plot(angles, vals, label=f"군집 {cid}", color=PALETTE[idx % len(PALETTE)], linewidth=1.8)
        ax.fill(angles, vals, color=PALETTE[idx % len(PALETTE)], alpha=0.08)
    ax.set_xticks(angles[:-1])
    ax.set_xticklabels(labels, fontsize=8)
    ax.set_title("군집별 지표 프로파일 (z-score)", fontsize=12, pad=20)
    ax.legend(loc="upper right", bbox_to_anchor=(1.25, 1.1), fontsize=8)
    fig.tight_layout()
    p = os.path.join(IMG_DIR, "radar.png")
    fig.savefig(p)
    plt.close(fig)
    paths["radar"] = p

    return paths


# ---------------------------------------------------------------- styling --

def style_header_row(ws, row, ncols, fill_hex=NAVY, font_color="FFFFFF"):
    fill = PatternFill("solid", fgColor=fill_hex)
    font = Font(bold=True, color=font_color, size=10)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
    ws.row_dimensions[row].height = 24


def thin_border():
    side = Side(style="thin", color="BFBFBF")
    return Border(left=side, right=side, top=side, bottom=side)


def autosize_columns(ws, min_width=10, max_width=40):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = len(str(cell.value))
            col = cell.column_letter
            widths[col] = max(widths.get(col, 0), length)
    for col, w in widths.items():
        ws.column_dimensions[col].width = min(max(w + 2, min_width), max_width)


def build_workbook(g1, cluster_mean, cluster_median, zscore, cluster_names, best_k, sil_scores, chart_paths):
    wb = Workbook()

    # ---------------- Summary sheet ----------------
    ws = wb.active
    ws.title = "요약"
    ws.sheet_view.showGridLines = False

    ws["B2"] = "GA 1그룹 지사 성과지표 군집분석 보고서"
    ws["B2"].font = Font(bold=True, size=18, color=NAVY)
    ws["B3"] = "성과지표 기반 유사 지사 그룹화 및 특징 분석"
    ws["B3"].font = Font(size=11, color=GREY)

    ws["B5"] = "분석 개요"
    ws["B5"].font = Font(bold=True, size=13, color=NAVY)
    overview = [
        ("분석 대상", f"그룹분류 = 1그룹 지사 {len(g1)}개"),
        ("사용 지표", f"성과지표 {len(METRICS)}개 (보장성NCEV, 가동5만/20만, 실수금률, 장기손해율 3종, 유지율 4종)"),
        ("분석 방법", "지표 극단치 완화(winsorize) 후 표준화 → K-means 군집분석"),
        ("군집 수(k)", f"{best_k}개 (실루엣 스코어 기준 최적값 {sil_scores[best_k]:.3f})"),
        ("Key", " / ".join(KEY_COLS)),
    ]
    r = 6
    for label, val in overview:
        ws.cell(row=r, column=2, value=label).font = Font(bold=True, size=10)
        ws.cell(row=r, column=3, value=val).font = Font(size=10)
        r += 1

    r += 1
    ws.cell(row=r, column=2, value="군집별 요약").font = Font(bold=True, size=13, color=NAVY)
    r += 1
    header_row = r
    headers = ["군집", "지사 수", "비중", "주요 특징"]
    for i, h in enumerate(headers):
        ws.cell(row=header_row, column=2 + i, value=h)
    style_header_row(ws, header_row, len(headers) + 1)
    # shift because header started at col B(2); style_header_row styles cols 1..ncols so redo properly
    for c in range(1, 2):
        ws.cell(row=header_row, column=c).fill = PatternFill(fill_type=None)
        ws.cell(row=header_row, column=c).font = Font()

    sizes = g1["군집"].value_counts().sort_index()
    total = len(g1)
    r += 1
    for cid in sorted(cluster_mean.index):
        ws.cell(row=r, column=2, value=f"군집 {cid}")
        ws.cell(row=r, column=3, value=int(sizes[cid]))
        ws.cell(row=r, column=4, value=round(sizes[cid] / total, 3))
        ws.cell(row=r, column=4).number_format = "0.0%"
        ws.cell(row=r, column=5, value=cluster_names[cid])
        fill = PatternFill("solid", fgColor=LIGHT_BLUE if cid % 2 == 0 else "FFFFFF")
        for c in range(2, 6):
            ws.cell(row=r, column=c).fill = fill
            ws.cell(row=r, column=c).border = thin_border()
            ws.cell(row=r, column=c).alignment = Alignment(vertical="center")
        r += 1

    r += 2
    ws.cell(row=r, column=2, value="핵심 인사이트").font = Font(bold=True, size=13, color=NAVY)
    r += 1
    for cid in sorted(cluster_mean.index):
        bullet = f"• 군집 {cid} ({sizes[cid]}개 지사, {sizes[cid]/total:.1%}): {cluster_names[cid]} 특성을 보임"
        ws.cell(row=r, column=2, value=bullet).font = Font(size=10)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        r += 1

    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 45
    ws.column_dimensions["F"].width = 20
    ws.freeze_panes = None

    # size chart embedded on summary
    img = XLImage(chart_paths["size"])
    img.width, img.height = 380, 240
    ws.add_image(img, "H6")

    # ---------------- Cluster Profile sheet ----------------
    ws2 = wb.create_sheet("군집 프로파일")
    ws2.sheet_view.showGridLines = False
    ws2["B2"] = "군집별 지표 평균 (원본 단위)"
    ws2["B2"].font = Font(bold=True, size=13, color=NAVY)

    start_row = 4
    headers = ["군집"] + [m.replace("_가상", "") for m in METRICS]
    for i, h in enumerate(headers):
        ws2.cell(row=start_row, column=2 + i, value=h)
    style_header_row(ws2, start_row, len(headers))

    for i, cid in enumerate(sorted(cluster_mean.index)):
        row = start_row + 1 + i
        ws2.cell(row=row, column=2, value=f"군집 {cid}")
        ws2.cell(row=row, column=2).font = Font(bold=True)
        for j, m in enumerate(METRICS):
            cell = ws2.cell(row=row, column=3 + j, value=round(float(cluster_mean.loc[cid, m]), 2))
            cell.border = thin_border()
            cell.number_format = "0.00"
        ws2.cell(row=row, column=2).border = thin_border()

    end_row = start_row + len(cluster_mean)
    end_col = 2 + len(METRICS)
    color_rule = ColorScaleRule(
        start_type="min", start_color="F8696B",
        mid_type="percentile", mid_value=50, mid_color="FFEB84",
        end_type="max", end_color="63BE7B",
    )
    ws2.conditional_formatting.add(
        f"{get_column_letter(3)}{start_row+1}:{get_column_letter(end_col)}{end_row}",
        color_rule,
    )

    r2 = end_row + 3
    ws2.cell(row=r2, column=2, value="군집별 지표 z-score (전체 평균 대비, 히트맵)").font = Font(bold=True, size=13, color=NAVY)
    img2 = XLImage(chart_paths["heatmap"])
    img2.width, img2.height = 620, 280
    ws2.add_image(img2, f"B{r2 + 2}")

    r3 = r2 + 18
    ws2.cell(row=r3, column=2, value="군집별 프로파일 레이더 차트").font = Font(bold=True, size=13, color=NAVY)
    img3 = XLImage(chart_paths["radar"])
    img3.width, img3.height = 420, 420
    ws2.add_image(img3, f"B{r3 + 2}")

    img4 = XLImage(chart_paths["silhouette"])
    img4.width, img4.height = 380, 240
    ws2.add_image(img4, f"J{r3 + 2}")

    ws2.column_dimensions["A"].width = 2
    for c in range(2, end_col + 1):
        ws2.column_dimensions[get_column_letter(c)].width = 14

    # ---------------- Branch Detail sheet ----------------
    ws3 = wb.create_sheet("지사별 상세")
    ws3.sheet_view.showGridLines = False
    cols = KEY_COLS + ["그룹분류", "군집"] + METRICS
    for i, h in enumerate(cols):
        ws3.cell(row=1, column=1 + i, value=h)
    style_header_row(ws3, 1, len(cols))

    detail = g1[cols].sort_values("군집").reset_index(drop=True)
    for r_i, row in enumerate(detail.itertuples(index=False), start=2):
        for c_i, val in enumerate(row, start=1):
            cell = ws3.cell(row=r_i, column=c_i, value=val)
            cell.border = thin_border()
            if cols[c_i - 1] in METRICS:
                cell.number_format = "0.00"

    n_rows = len(detail) + 1
    n_cols = len(cols)
    ws3.freeze_panes = "E2"
    ws3.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{n_rows}"

    cluster_col = cols.index("군집") + 1
    # 군집별 밴딩 색상 (연한 배경)
    band_colors = ["EAF1FA", "FBEEE6", "EDF3E8", "F1E9F5", "FBF3DE"]
    for r_i in range(2, n_rows + 1):
        cid_val = ws3.cell(row=r_i, column=cluster_col).value
        fill = PatternFill("solid", fgColor=band_colors[int(cid_val) % len(band_colors)])
        for c_i in range(1, n_cols + 1):
            ws3.cell(row=r_i, column=c_i).fill = fill

    autosize_columns(ws3, min_width=9, max_width=22)

    # ---------------- Methodology sheet ----------------
    ws4 = wb.create_sheet("방법론")
    ws4.sheet_view.showGridLines = False
    ws4["B2"] = "분석 방법론"
    ws4["B2"].font = Font(bold=True, size=14, color=NAVY)
    lines = [
        "1. 데이터: 업로드된 GA 지사 성과 데이터 중 그룹분류='1그룹' 지사 472개를 분석 대상으로 필터링.",
        "2. 전처리: 11개 성과지표 각각에 대해 1~99 percentile로 winsorize(극단치 완화) 후 StandardScaler로 표준화.",
        "   - 보장성NCEV, 가동5만, 가동20만 지표는 서로 상관 0.82~0.94로 높으나, 각 지표가 지닌 의미를 보존하기 위해",
        "     별도 차원 축소 없이 극단치 처리만 수행하여 모두 군집화에 사용.",
        f"3. 군집화: K-means, k=3~5 후보 중 실루엣 스코어가 가장 높은 k={best_k} 선정"
        f" (스코어: {', '.join(f'k={k}:{v:.3f}' for k, v in sil_scores.items())}).",
        "4. 군집 해석: 군집별 지표 평균의 전체 평균 대비 z-score를 산출하여, 절대값이 큰 상위 2개 지표로 군집 특징을 명명.",
        "5. 시각화: 군집별 z-score 히트맵, 레이더 차트, 군집 크기 막대그래프, k 선정 근거(실루엣 스코어) 막대그래프.",
        "6. 검증: 필터링 행 수(472), key 중복 여부, 결측치 여부, 최소 군집 크기(>=10)를 스크립트 내에서 자동 검증.",
    ]
    for i, line in enumerate(lines):
        ws4.cell(row=4 + i, column=2, value=line).font = Font(size=10)
        ws4.merge_cells(start_row=4 + i, start_column=2, end_row=4 + i, end_column=10)
    ws4.column_dimensions["A"].width = 2
    ws4.column_dimensions["B"].width = 14

    wb.save(XLSX_PATH)


def main():
    g1 = load_group1()
    g1, best_k, sil_scores = cluster(g1)
    cluster_mean, cluster_median, zscore = build_profile(g1)
    cluster_names = name_clusters(zscore)
    chart_paths = make_charts(g1, cluster_mean, zscore, best_k, sil_scores)
    build_workbook(g1, cluster_mean, cluster_median, zscore, cluster_names, best_k, sil_scores, chart_paths)

    # self-verify
    from openpyxl import load_workbook
    wb = load_workbook(XLSX_PATH)
    assert set(wb.sheetnames) == {"요약", "군집 프로파일", "지사별 상세", "방법론"}
    ws3 = wb["지사별 상세"]
    assert ws3.max_row - 1 == len(g1) == 472
    cluster_col_idx = KEY_COLS.__len__() + 2
    missing = sum(1 for r in range(2, ws3.max_row + 1) if ws3.cell(row=r, column=cluster_col_idx).value is None)
    assert missing == 0, "군집 라벨 결측 존재"
    print(f"검증 완료: {ws3.max_row - 1}행, 시트 {wb.sheetnames}")
    print(f"엑셀 저장 완료: {XLSX_PATH}")


if __name__ == "__main__":
    main()
